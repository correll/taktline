#!/usr/bin/env python3
"""The second tier: a language model reviews the what-if model's answers.

The line doctor's rules and the what-if model are heuristics. Each is right to
the extent that the line resembles the idealisation it encodes, and an
experienced engineer discounts an answer when the record shows that the line
works differently: a station that never has room downstream, two machines that
are never busy at once. This tool asks a language model to take that role. For
every what-if question it states how likely the answer is to be wrong for this
line, and why.

It is **optional and off unless you run it**: nothing in Taktline calls it, and
it is the one part of Taktline that sends the plant's record to a third party.
It talks to a running server over the public HTTP API, signed in as a member,
exactly as the browser does, and calls Claude through the Claude Code command
line with the caller's own credentials (`claude` must be installed and signed
in). The first tier, statistics, line doctor and what-if model, needs none of
this and runs on the plant's own computer.

    .venv/bin/python review_whatif.py --email you@example.com --company "My line" --auto
    .venv/bin/python review_whatif.py --email you@example.com --scenarios questions.json --out review.json

`--auto` asks the four questions the study asked of every line: maintenance at
the least reliable station, a second server at the most loaded one, SMED at the
station with the most changeover time, and 10% more demand. `--scenarios` takes
a JSON list in the body format of `POST /analysis/what-if`.

The instructions, the reply format and the input are those evaluated in the
paper (Correll, "Digital value stream mapping from operator-recorded events"):
`PROMPT_SHA256` pins the instructions, and `tests/test_review_whatif.py` fails
if they drift. What the reviewer is given: the line's stations and routes,
Taktline's statistics, the line doctor's findings, the what-if answers with the
model's baseline and the observed lead time, and evidence from the record of
sessions and moves over the last 30 days -- how often each pair of stations works
at once against independent working, how long lots queue, how long finished
lots wait before they move on. Never operator names, order data or tracking
links.
"""
from __future__ import annotations

import argparse
import bisect
import hashlib
import io
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta, timezone

import httpx

MODEL = "claude-opus-5-5"
WINDOW_DAYS = 30

SYSTEM = 'You review the output of a production-line simulation model for a manufacturing engineer. You are careful, quantitative and candid, and you answer only in the JSON format requested.'

TASK = 'A small factory tracks its line with Taktline, a digital kanban system: operators press Start, Pause and Stop on lots at each station and record every move between stations. From that record Taktline computes statistics, runs a rule base of lean diagnostics (the "line doctor"), and answers what-if questions with a discrete-event model.\n\nHow the what-if model works. It replays the recorded period: every lot that left an entry station enters the model at the same moment and quantity. Work per unit is resampled per station and value stream from the recorded work sessions (including the interruptions those sessions contained). Breakdowns occur at the recorded rate per working hour with the recorded repair times; changeovers occur when the product changes, with the recorded probability and durations. Each station serves its queue first in, first out with its own number of servers; buffers between stations are unlimited, and each station is assumed to be staffed independently of the others. Lots follow the routes of their value stream; a merge waits for a matched set. A scenario changes one thing (removes a station\'s breakdowns or changeovers, adds a server, speeds a station up, or advances every release to raise demand) and is simulated against the baseline with common random numbers; the answer is the predicted reduction in mean lead time with a 95% interval over the model\'s own replications. The model\'s baseline lead time is shown next to the lead time actually observed, as a check of its fit.\n\nThe model is a heuristic: its assumptions hold on some lines and not on others. An experienced engineer looking at the whole record can often see that a line works differently from what the model assumes, and that its answer is therefore simply wrong for this line. Below is everything Taktline recorded and computed for one line over about eight working days, including evidence from the record itself (how lots queue and wait, and which stations work at the same time). For EACH what-if question, estimate the probability that the model\'s answer is wrong for this line: that the true reduction in mean lead time, averaged over many repetitions of a period like this one, lies OUTSIDE 0.80 to 1.25 times the predicted reduction. Do not count the chance variation of a single period; judge whether the model fits the line. Look for behaviour the model cannot represent, compare the model\'s baseline with what was observed, and consider how close stations run to saturation. Give a probability for every question and one sentence of reason.\n\nDATA (JSON):\n{data}'

SCHEMA = {'type': 'object', 'properties': {'answers': {'type': 'array', 'items': {'type': 'object', 'properties': {'question': {'type': 'string'}, 'p_outside': {'type': 'number', 'minimum': 0, 'maximum': 1}, 'reason': {'type': 'string'}}, 'required': ['question', 'p_outside', 'reason']}}}, 'required': ['answers']}


PROMPT_SHA256 = "3e42be06ef3f0079618ae22780ddd0ceba7e8c272f542ffee29fd7a0107bc4d4"


def prompt_sha256() -> str:
    return hashlib.sha256((SYSTEM + "\n" + TASK + "\n" + json.dumps(SCHEMA, sort_keys=True)).encode()).hexdigest()


def prompt_text(data: dict) -> str:
    return TASK.format(data=json.dumps(data, separators=(",", ":")))


# ── HTTP ──────────────────────────────────────────────────────────────────────

class Api:
    """A signed-in member's session. `client` may be any httpx.Client (the
    tests pass FastAPI's TestClient)."""

    def __init__(self, client: httpx.Client, company_id: str | None = None):
        self.client = client
        self.company_id = company_id

    def req(self, method: str, path: str, body=None) -> httpx.Response:
        h = {}
        csrf = self.client.cookies.get("taktline_csrf")
        if csrf and method.upper() != "GET":
            h["X-CSRF-Token"] = csrf
        if self.company_id:
            h["X-Company-Id"] = self.company_id
        return self.client.request(method, path, json=body, headers=h)

    def get(self, path: str):
        r = self.req("GET", path)
        r.raise_for_status()
        return r.json()


def login(api: Api, email: str, password: str, company: str | None) -> str:
    r = api.req("POST", "/auth/login", {"email": email, "password": password})
    if r.status_code != 200:
        sys.exit(f"login failed: {r.status_code} {r.text}")
    companies = r.json().get("companies", [])
    if company:
        match = [c for c in companies if c["name"].lower() == company.lower()]
        if not match:
            sys.exit(f"no company named {company!r}")
        chosen = match[0]
    elif len(companies) == 1:
        chosen = companies[0]
    else:
        sys.exit("several companies -- pick one with --company")
    api.company_id = chosen.get("companyId") or chosen.get("id")
    return chosen["name"]


# ── what Taktline recorded and computed ─────────────────────────────────────

def line_structure(stations: list, routes: list, processes: list, worked: set = frozenset()) -> dict:
    """Stations with a role and each value stream's routes. A station bound
    to an integration partner is a boundary, not part of the line. An entry
    is where lots originate and nobody works; lots are often created at the
    first work station itself, which is then a work station."""
    name = {s["_id"]: s["name"] for s in stations}
    partner = {s["_id"] for s in stations if s.get("integrationProvider")}
    pname = {p["_id"]: p["name"] for p in processes}
    flow = [r for r in routes if r["load"] not in partner and r["unload"] not in partner]
    incoming = {r["unload"] for r in flow}
    role = {}
    for s in stations:
        if s["_id"] in partner:
            continue
        role[s["name"]] = ("storage" if s.get("isStorage") else
                           "entry" if s["_id"] not in incoming and s["_id"] not in worked else "work")
    vs: dict = {}
    for r in flow:
        vs.setdefault(pname.get(r["processId"], r["processId"]), []).append([name[r["load"]], name[r["unload"]]])
    return {"stations": [{"name": n, "role": x} for n, x in role.items()], "value_streams": vs}


def _r(x, n=3):
    return round(x, n) if isinstance(x, float) else x


def statistics(stats: list, name: dict) -> list:
    keep = ("sessionCount", "totalUnits", "cycleTimeSec", "lineCycleTimeSec", "oee", "uptime",
            "downtimeSec", "breakdowns", "changeoverSec", "changeovers", "wipAvg", "unstoppedSessions")
    out = []
    for s in stats:
        if s["stationId"] not in name:
            continue
        row = {"station": name[s["stationId"]]}
        row.update({k: _r(s.get(k)) for k in keep})
        row["shareOfTimeWorked"] = row.pop("oee")
        out.append(row)
    return out


def findings(report: dict | None, name: dict) -> list:
    out = []
    for f in (report or {}).get("findings", []):
        if f["rule"] in ("single-point-of-failure", "no-takt", "long-transport"):
            continue                      # topology notes with no bearing on the model's fit
        out.append({"rule": f["rule"], "title": f.get("title"),
                    "stations": [name.get(x, x) for x in f.get("stations") or []],
                    "metrics": {k: _r(v) for k, v in (f.get("metrics") or {}).items()},
                    "detail": f.get("detail")})
    return out


def auto_scenarios(stats: list, name: dict, work: set) -> list:
    """The study's rule: maintenance where uptime is lowest, a second server
    where the share of time worked is highest, SMED where changeovers cost
    most, and 10% more demand."""
    rows = [s for s in stats if name.get(s["stationId"]) in work]
    out = []
    down = [s for s in rows if s.get("uptime") is not None and s["uptime"] < 0.999]
    if down:
        s = min(down, key=lambda s: s["uptime"])
        out.append({"name": f"{name[s['stationId']]}: no breakdowns (TPM)",
                    "changes": [{"stationId": s["stationId"], "breakdowns": False}]})
    busy = [s for s in rows if s.get("oee")]
    if busy:
        s = max(busy, key=lambda s: s["oee"])
        out.append({"name": f"{name[s['stationId']]}: second server",
                    "changes": [{"stationId": s["stationId"], "servers": 2}]})
    co = [s for s in rows if (s.get("changeoverSec") or 0) > 0]
    if co:
        s = max(co, key=lambda s: s["changeoverSec"])
        out.append({"name": f"{name[s['stationId']]}: no changeovers (SMED)",
                    "changes": [{"stationId": s["stationId"], "changeovers": False}]})
    out.append({"name": "Demand +10%", "changes": [], "demandFactor": 1.1})
    return out


def what_if(api: Api, scenarios: list, name: dict) -> dict:
    r = api.req("POST", "/analysis/what-if", {"scenarios": scenarios})
    r.raise_for_status()
    w = r.json()
    base = w["baseline"]
    stations = {name.get(k, k): v for k, v in (base.get("stations") or {}).items()}
    return {"model_baseline": {"leadTimeHours": base["leadTimeHours"], "wipUnits": base["wipUnits"],
                               "leadSpreadHours": base.get("leadSpreadHours"), "stations": stations},
            "fit_check": w["calibration"],
            "questions": [{"question": s["name"], "predictedReductionHours": s["gainHours"],
                           "ci95Hours": s["gainCi95Hours"],
                           "predictedLeadTimeHours": s["leadTimeHours"]} for s in w["scenarios"]]}


# ── evidence from the record of sessions and moves ──────────────────────────

def _utc(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v if v.tzinfo else v.replace(tzinfo=timezone.utc)
    t = datetime.fromisoformat(str(v).replace("Z", "+00:00"))
    return t if t.tzinfo else t.replace(tzinfo=timezone.utc)


def read_export(content: bytes) -> tuple[list, list]:
    """Sessions and lot events from GET /reports/export.xlsx."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True)

    def rows(sheet):
        it = wb[sheet].iter_rows(values_only=True)
        head = list(next(it))
        return [dict(zip(head, r)) for r in it]
    sessions = [{"start": _utc(r["Started (UTC)"]), "end": _utc(r["Ended (UTC)"]), "station": r["Station"],
                 "lot": r.get("Lot id") or ""} for r in rows("Work sessions")]
    events = [{"t": _utc(r["When (UTC)"]), "type": r["Event"], "from": r["From"] or None, "to": r["To"] or None,
               "note": r["Note"] or "", "lot": r["Lot id"]} for r in rows("Lot history")]
    return sessions, events


class OpenClock:
    """Open minutes between two moments, from the plant's calendar."""

    def __init__(self, intervals: list):
        self.iv = [(a, b) for a, b in intervals if b > a]
        self.starts = [a for a, _ in self.iv]
        self.cum = [0.0]
        for a, b in self.iv:
            self.cum.append(self.cum[-1] + (b - a).total_seconds() / 60)

    def at(self, t: datetime) -> float:
        i = bisect.bisect_right(self.starts, t) - 1
        if i < 0:
            return 0.0
        a, b = self.iv[i]
        return self.cum[i] + (min(t, b) - a).total_seconds() / 60

    def between(self, x: datetime, y: datetime) -> float:
        return max(0.0, self.at(y) - self.at(x))

    @property
    def total(self) -> float:
        return self.cum[-1]


def _union(iv):
    out = []
    for a, b in sorted(iv):
        if out and a <= out[-1][1]:
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return out


def _clip(iv, opn):
    out = []
    for a, b in iv:
        for lo, hi in opn:
            x, y = max(a, lo), min(b, hi)
            if y > x:
                out.append([x, y])
    return out


def _overlap(x, y):
    i = j = 0
    tot = 0.0
    while i < len(x) and j < len(y):
        lo, hi = max(x[i][0], y[j][0]), min(x[i][1], y[j][1])
        if hi > lo:
            tot += (hi - lo).total_seconds() / 60
        if x[i][1] < y[j][1]:
            i += 1
        else:
            j += 1
    return tot


def record_evidence(sessions: list, events: list, work: list, intervals: list,
                    since: datetime, until: datetime) -> dict:
    """What an engineer would read off the record beyond per-station averages,
    in the plant's open time over [since, until]."""
    clock = OpenClock(intervals)
    open_min = clock.total or 1.0
    busy = {s: [] for s in work}
    first_start, last_end = {}, {}
    for s in sorted(sessions, key=lambda s: s["start"]):
        if s["station"] not in busy or s["start"] is None:
            continue
        end = s["end"] or until
        busy[s["station"]].append([s["start"], end])
        k = (s["lot"], s["station"])
        first_start.setdefault(k, s["start"])
        if s["end"] is not None:
            last_end[k] = s["end"]
    arrive, leave, merged = {}, {}, {}
    queue_steps = {s: [] for s in work}
    for e in sorted(events, key=lambda e: e["t"]):
        if e["type"] == "moved" and e["from"] and e["to"]:
            arrive.setdefault((e["lot"], e["to"]), e["t"])
            leave[(e["lot"], e["from"])] = e["t"]
            if e["to"] in queue_steps:
                queue_steps[e["to"]].append((e["t"], +1))
        elif e["type"] == "updated" and e["from"] and str(e["note"]).startswith("merged"):
            merged.setdefault((e["lot"], e["from"]), e["t"])
    for k, t in list(first_start.items()) + [(k, t) for k, t in merged.items() if k not in first_start]:
        if k in arrive and k[1] in queue_steps:
            queue_steps[k[1]].append((t, -1))
    opn = [[a, b] for a, b in intervals]
    U = {s: _clip(_union(busy[s]), opn) for s in work}
    frac = {s: sum((b - a).total_seconds() / 60 for a, b in U[s]) / open_min for s in work}
    pairs = []
    for i, a in enumerate(work):
        for b in work[i + 1:]:
            if frac[a] < 0.05 or frac[b] < 0.05:
                continue
            pairs.append({"stations": [a, b], "bothWorkingShare": round(_overlap(U[a], U[b]) / open_min, 3),
                          "ifIndependent": round(frac[a] * frac[b], 3)})

    def q(v):
        if not v:
            return None
        v = sorted(v)
        p90 = v[min(len(v) - 1, int(round(0.9 * (len(v) - 1))))]
        med = v[len(v) // 2] if len(v) % 2 else (v[len(v) // 2 - 1] + v[len(v) // 2]) / 2
        return {"median_min": round(med, 1), "p90_min": round(p90, 1)}
    stations = {}
    for s in work:
        qw = [clock.between(arrive[k], first_start[k]) for k in first_start
              if k[1] == s and k in arrive and arrive[k] >= since]
        hold = [clock.between(last_end[k], leave[k]) for k in last_end
                if k[1] == s and k in leave and last_end[k] >= since]
        lvl, last, area, empty, mx = 0, since, 0.0, 0.0, 0
        for t, d in sorted(queue_steps[s], key=lambda x: (x[0], -x[1])):
            if t > since:
                ow = clock.between(last, t)
                area += lvl * ow
                empty += ow if lvl == 0 else 0.0
                last = t
            lvl = max(0, lvl + d)
            if t >= since:
                mx = max(mx, lvl)
        ow = clock.between(last, until)
        area += lvl * ow
        empty += ow if lvl == 0 else 0.0
        stations[s] = {"workingShare": round(frac[s], 3),
                       "lotsQueuedBeforeStart": q(qw),
                       "finishedLotsWaitBeforeMoving": q(hold),
                       "queueLots": {"timeAverage": round(area / open_min, 2), "max": mx,
                                     "shareOfTimeEmpty": round(empty / open_min, 3)}}
    return {"note": "open time of the plant's calendar, measurement window only; minutes; "
                    "queue = lots arrived at the station but not yet started (or merged)",
            "stations": stations, "pairsWorkingAtOnce": pairs}


def build_input(api: Api, scenarios: list | None, now: datetime | None = None) -> tuple[dict, list]:
    """Everything the reviewer sees, gathered over the public API."""
    from app.core.schedule import Schedule, open_intervals
    stations = api.get("/stations/")
    name = {s["_id"]: s["name"] for s in stations}
    stats = api.get("/sessions/stats")
    worked = {s["stationId"] for s in stats if s.get("sessionCount")}
    line = line_structure(stations, api.get("/routes/"), api.get("/processes/"), worked)
    work = [x["name"] for x in line["stations"] if x["role"] == "work"]
    r = api.req("GET", "/analysis/")
    report = r.json() if r.status_code == 200 else None     # the line doctor may not be on this plan
    if not scenarios:
        scenarios = auto_scenarios(stats, name, set(work))
    gs = api.get("/global-settings/")
    r = api.req("GET", "/reports/export.xlsx")
    r.raise_for_status()
    sessions, events = read_export(r.content)
    until = now or max([e["t"] for e in events] + [s["end"] or s["start"] for s in sessions]
                       + [datetime.now(timezone.utc)])
    since = until - timedelta(days=WINDOW_DAYS)
    first = min([e["t"] for e in events if e["t"]] + [until])
    schedule = Schedule(gs.get("workHours"), gs.get("breaks") or [])
    intervals = open_intervals(max(since, first), until, schedule, gs.get("timezone") or "UTC")
    data = {"line": line, "statistics": statistics(stats, name),
            "record_evidence": record_evidence(sessions, events, work, intervals, max(since, first), until),
            "line_doctor": findings(report, name),
            "what_if": what_if(api, scenarios, name)}
    return data, [q["question"] for q in data["what_if"]["questions"]]


# ── the reviewer ─────────────────────────────────────────────────────────────

def ask_claude(data: dict, model: str = MODEL) -> dict:
    """One headless Claude Code call with every tool and customisation
    disabled, from an empty directory, so nothing but the input reaches it."""
    with tempfile.TemporaryDirectory() as empty:
        r = subprocess.run(
            ["claude", "-p", "--model", model, "--safe-mode", "--tools", "", "--no-session-persistence",
             "--system-prompt", SYSTEM, "--output-format", "json", "--json-schema", json.dumps(SCHEMA)],
            input=prompt_text(data), capture_output=True, text=True, cwd=empty, timeout=1800)
    if r.returncode != 0:
        raise RuntimeError(f"claude exited {r.returncode}: {r.stderr[-500:]}")
    env = json.loads(r.stdout)
    body = env.get("structured_output") or json.loads(env.get("result", "{}"))
    return {"answers": body["answers"], "model_usage": env.get("modelUsage"),
            "cost_usd": env.get("total_cost_usd")}


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--url", default=os.environ.get("TAKTLINE_URL", "http://127.0.0.1:8000"))
    ap.add_argument("--email", required=True)
    ap.add_argument("--password", default=None, help="or $TAKTLINE_PASSWORD")
    ap.add_argument("--company", default=None)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--auto", action="store_true", help="the study's four questions")
    g.add_argument("--scenarios", help="JSON file: a list of what-if scenarios")
    ap.add_argument("--model", default=MODEL)
    ap.add_argument("--out", default=None, help="write input and replies here (JSON)")
    ap.add_argument("--dry-run", action="store_true", help="build and print the input; send nothing")
    a = ap.parse_args()
    if prompt_sha256() != PROMPT_SHA256:
        sys.exit("the reviewer's instructions differ from the evaluated ones; refusing to run")
    password = a.password or os.environ.get("TAKTLINE_PASSWORD")
    if not password:
        sys.exit("--password or $TAKTLINE_PASSWORD is required")
    api = Api(httpx.Client(base_url=a.url.rstrip("/"), timeout=600.0, follow_redirects=True))
    company = login(api, a.email, password, a.company)
    scenarios = json.loads(open(a.scenarios).read()) if a.scenarios else None
    data, questions = build_input(api, scenarios)
    if a.dry_run:
        print(prompt_text(data))
        return
    print(f"Reviewing {len(questions)} what-if answers for {company} with {a.model} ...", file=sys.stderr)
    res = ask_claude(data, a.model)
    for x in res["answers"]:
        print(f"{x['p_outside']:5.2f}  {x['question']}\n       {x['reason']}")
    if a.out:
        with open(a.out, "w") as f:
            json.dump({"input": data, "prompt_sha256": PROMPT_SHA256, "model": a.model, **res}, f, indent=1)


if __name__ == "__main__":
    main()
