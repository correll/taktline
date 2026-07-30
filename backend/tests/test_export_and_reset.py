"""The spreadsheet export and the statistics wipe.

The export is checked by **opening the workbook** and reading cells, not by
asserting a 200 — a corrupt xlsx is still a 200, and the whole point of the
format over CSV is that Excel opens it without an import dialog. So: real
sheets, real typed cells (a timestamp must arrive as a datetime and a duration
as a number, or the receiver is back to parsing strings), and the numbers must
agree with what the Statistics page reports.

The wipe is checked for what it *keeps* as much as what it removes: clearing
measurements must never be a way to lose a customer's line.
"""

from io import BytesIO

from openpyxl import load_workbook

from tests.helpers import Actor, bootstrap_admin

STATION = {"description": "", "adapterType": None, "adapterConfig": {}, "color": None,
           "icon": None, "position": {"x": 1, "y": 1, "width": 60, "height": 60}}


def _line(a: Actor):
    st = lambda n, **kw: a.req("POST", "/stations/", {**STATION, "name": n, **kw}).json()["_id"]
    one, two = st("Cutting"), st("Welding")
    proc = a.req("POST", "/processes/", {"name": "Frames"}).json()["_id"]
    a.req("POST", "/routes/", {"processId": proc, "load": one, "unload": two, "quantity": 1})
    tmpl = a.req("POST", "/lot-templates/", {"name": "Frame", "processId": proc,
                                             "startStationId": one}).json()["_id"]
    card = a.req("POST", "/cards/", {"lotTemplateId": tmpl, "processId": proc, "lotNum": "0042",
                                     "totalQuantity": 6, "bins": {one: {"count": 6}}}).json()
    s = a.req("POST", "/sessions/", {"stationId": one, "cardId": card["_id"],
                                     "operatorId": "tablet-01", "quantity": 6})
    a.req("POST", f"/sessions/{s.json()['_id']}/end", {"quantity": 6})
    a.req("POST", f"/cards/{card['_id']}/advance",
          {"fromStationId": one, "toStationId": two, "quantity": 3})
    a.req("POST", f"/cards/{card['_id']}/advance",
          {"fromStationId": one, "toStationId": two, "quantity": 3})
    return one, two, card


def _workbook(a: Actor, query: str = ""):
    r = a.req("GET", f"/reports/export.xlsx{query}")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"]
    assert ".xlsx" in r.headers["content-disposition"]
    return load_workbook(BytesIO(r.content))


def test_the_workbook_opens_and_carries_every_table(actor):
    a = actor(); bootstrap_admin(a)
    _line(a)
    wb = _workbook(a)
    assert wb.sheetnames == ["About", "Stations", "Work sessions",
                             "Lot history", "Lots now", "Orders"]

    st = wb["Stations"]
    assert [c.value for c in st[1]] == [
        "Station", "Process(es)", "Storage", "Sessions", "Units", "Working time (s)",
        "Cycle time (s/piece)", "Line cycle time (s/piece)", "OEE"]
    by_name = {r[0]: r for r in st.iter_rows(min_row=2, values_only=True)}
    assert by_name["Cutting"][3] == 1                 # one session
    assert by_name["Cutting"][4] == 6                 # six units
    # Cycle time must be a NUMBER, not "0.1 s" — the receiver charts it.
    assert isinstance(by_name["Cutting"][6], (int, float))
    # OEE is stored as a fraction and *formatted* as a percentage, so the cell
    # is still arithmetic. A string "9.1%" would be useless in a pivot table.
    oee_cell = st.cell(row=[r[0] for r in st.iter_rows(min_row=2, values_only=True)]
                       .index("Cutting") + 2, column=9)
    assert oee_cell.number_format == "0.0%"
    assert oee_cell.value is None or isinstance(oee_cell.value, (int, float))


def test_timestamps_arrive_as_dates_not_text(actor):
    """The reason for xlsx over CSV: Excel guesses at CSV text and mangles it.
    A datetime cell removes the guess."""
    import datetime as dt
    a = actor(); bootstrap_admin(a)
    _line(a)
    wb = _workbook(a)
    started = wb["Work sessions"].cell(row=2, column=1).value
    assert isinstance(started, dt.datetime), f"got {started!r}"
    assert started.tzinfo is None, "Excel stores wall clock; the header says UTC"

    when = wb["Lot history"].cell(row=2, column=1).value
    assert isinstance(when, dt.datetime)


def test_a_lot_code_that_looks_like_a_number_survives(actor):
    """'0042' through a CSV comes back as 42. It must not here."""
    a = actor(); bootstrap_admin(a)
    _line(a)
    wb = _workbook(a)
    codes = {r[0] for r in wb["Lots now"].iter_rows(min_row=2, values_only=True)}
    assert "0042" in codes, codes


def test_the_history_sheets_carry_the_movements(actor):
    a = actor(); bootstrap_admin(a)
    one, two, _ = _line(a)
    wb = _workbook(a)
    moves = [r for r in wb["Lot history"].iter_rows(min_row=2, values_only=True)
             if r[1] == "moved"]
    assert len(moves) == 2
    assert {r[4] for r in moves} == {"Cutting"} and {r[5] for r in moves} == {"Welding"}
    assert sum(r[6] for r in moves) == 6


def test_a_window_narrows_the_export_the_way_the_page_does(actor):
    a = actor(); bootstrap_admin(a)
    _line(a)
    wb = _workbook(a, "?window=1h")
    cutting = next(r for r in wb["Stations"].iter_rows(min_row=2, values_only=True)
                   if r[0] == "Cutting")
    assert cutting[3] == 1, "work done seconds ago belongs to the past hour"
    assert wb["About"]["B4"].value == "1h"


# ─── The wipe ────────────────────────────────────────────────────────────────

def test_resetting_statistics_clears_measurements_and_keeps_the_line(actor):
    a = actor(); bootstrap_admin(a)
    one, two, card = _line(a)
    assert a.req("GET", "/sessions/").json(), "precondition: sessions exist"

    r = a.req("POST", "/admin/reset-statistics")
    assert r.status_code == 200, r.text
    assert r.json()["workSessionsDeleted"] == 1
    assert r.json()["cardEventsDeleted"] >= 2

    assert a.req("GET", "/sessions/").json() == []
    assert a.req("GET", f"/cards/{card['_id']}/events").json() == []
    assert a.req("GET", "/sessions/stats").json() == []

    # The line and the work in progress are untouched — this clears the
    # measurement, not the factory.
    assert len(a.req("GET", "/stations/").json()) == 2
    assert len(a.req("GET", "/routes/").json()) == 1
    assert len(a.req("GET", "/lot-templates/").json()) == 1
    fresh = a.req("GET", f"/cards/{card['_id']}").json()
    assert fresh["bins"][two]["count"] == 6

    # And the line still records: the next session starts a clean history.
    s = a.req("POST", "/sessions/", {"stationId": two, "quantity": 1})
    a.req("POST", f"/sessions/{s.json()['_id']}/end", {"quantity": 1})
    assert len(a.req("GET", "/sessions/").json()) == 1


def test_only_a_super_admin_may_wipe_the_statistics(actor):
    a = actor(); bootstrap_admin(a)
    _line(a)
    from tests.helpers import add_member
    member = actor()
    add_member(a, member, role="admin")
    assert member.req("POST", "/admin/reset-statistics").status_code == 403
    assert len(a.req("GET", "/sessions/").json()) == 1, "the refused call changed nothing"


# ─── The detail behind a number ──────────────────────────────────────────────

def test_detail_returns_samples_not_more_averages(actor):
    """The point of the drill-down: an average hides the outlier that explains
    the station. So it must hand back the individual observations."""
    a = actor(); bootstrap_admin(a)
    one, _two, _card = _line(a)

    d = a.req("GET", f"/sessions/stats/{one}/detail").json()
    assert d["stationName"] == "Cutting"
    assert len(d["cycleSamples"]) == 1
    s = d["cycleSamples"][0]
    assert s["units"] == 6 and s["secPerPiece"] > 0
    # Two departures were recorded; the first has nothing to measure from, so
    # exactly one interval survives — the same rule line cycle time uses.
    assert len(d["departureSamples"]) == 1

    # Activity lands in a real hour, and the schedule comes along so a client
    # can show work happening outside it.
    worked = [h for h in d["byHour"] if h["sessions"]]
    assert len(worked) == 1 and worked[0]["units"] == 6
    assert len(d["byHour"]) == 24
    assert d["openHours"] == list(range(9, 17))       # the default 09:00–17:00


def test_detail_honours_the_window_and_refuses_a_stranger(actor):
    a = actor(); bootstrap_admin(a)
    one, _two, _card = _line(a)
    assert a.req("GET", f"/sessions/stats/{one}/detail?window=1h").json()["window"] == "1h"

    other = actor(); bootstrap_admin(other)
    assert other.req("GET", f"/sessions/stats/{one}/detail").status_code == 404
