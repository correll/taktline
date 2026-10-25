"""The second-tier reviewer (review_whatif.py).

What is pinned here is what the paper evaluated: the instructions are the
frozen ones, the evidence shows the two failure modes it was built to show,
the input is gathered from the record over the public API, and nothing
personal goes into it. The language model itself is not called: what it
answers is the study's business, not a unit test's.
"""
from datetime import datetime, timedelta, timezone
import io
import json

import pytest

import review_whatif as R
from tests.helpers import bootstrap_admin
from tests.test_stoppages import _lot, _line, _start, clock  # noqa: F401

T0 = datetime(2027, 1, 4, 6, 0, tzinfo=timezone.utc)


def _m(minutes):
    return T0 + timedelta(minutes=minutes)


def test_the_instructions_are_the_evaluated_ones():
    assert R.prompt_sha256() == R.PROMPT_SHA256 == \
        "3e42be06ef3f0079618ae22780ddd0ceba7e8c272f542ffee29fd7a0107bc4d4"


def test_two_stations_one_operator_never_work_at_once():
    """One person running two stations alternates between them: they never
    work at the same moment, although each is busy a good share of the time.
    Two independently staffed stations overlap as chance says."""
    shared = [{"start": _m(20 * k), "end": _m(20 * k + 10), "station": "A", "lot": f"a{k}"} for k in range(24)] + \
             [{"start": _m(20 * k + 10), "end": _m(20 * k + 20), "station": "B", "lot": f"b{k}"} for k in range(24)]
    ev = R.record_evidence(shared, [], ["A", "B"], [(T0, _m(480))], T0, _m(480))
    p = ev["pairsWorkingAtOnce"][0]
    assert p["bothWorkingShare"] == 0 and p["ifIndependent"] == pytest.approx(0.25, abs=0.01)
    free = [{"start": _m(20 * k), "end": _m(20 * k + 10), "station": "A", "lot": f"a{k}"} for k in range(24)] + \
           [{"start": _m(20 * k + 5), "end": _m(20 * k + 15), "station": "B", "lot": f"b{k}"} for k in range(24)]
    ev = R.record_evidence(free, [], ["A", "B"], [(T0, _m(480))], T0, _m(480))
    assert ev["pairsWorkingAtOnce"][0]["bothWorkingShare"] > 0.2


def test_a_blocked_station_holds_its_finished_lots():
    """A lot finished at A but moved on only when B has room waits at A: the
    wait between the end of its work and its move is the signature of
    blocking, and a free-flowing station shows none."""
    sessions, events = [], []
    for k in range(6):
        sessions.append({"start": _m(30 * k), "end": _m(30 * k + 10), "station": "A", "lot": f"L{k}"})
        events.append({"t": _m(30 * k + 25), "type": "moved", "from": "A", "to": "B", "note": "", "lot": f"L{k}"})
        sessions.append({"start": _m(30 * k + 25), "end": _m(30 * k + 30), "station": "B", "lot": f"L{k}"})
    ev = R.record_evidence(sessions, events, ["A", "B"], [(T0, _m(200))], T0, _m(200))
    assert ev["stations"]["A"]["finishedLotsWaitBeforeMoving"]["median_min"] == pytest.approx(15)
    assert ev["stations"]["B"]["lotsQueuedBeforeStart"]["median_min"] == 0


def test_the_reviewer_sees_the_line_the_record_and_the_answers(actor, clock, monkeypatch):  # noqa: F811
    """Built over the public API from a small recorded line: the stations
    with their roles, Taktline's statistics, evidence from the sessions and
    moves, and the what-if answers with the model's fit check. Operator ids
    and lot codes never reach the input."""
    a = actor(); bootstrap_admin(a)
    a.req("PUT", "/global-settings/", {"workHours": [[15] * 24 for _ in range(7)]})
    ids, tmpl = _line(a)
    for k in range(6):
        card = _lot(a, ids, tmpl)
        sid = _start(a, ids, card, op="ann-smith")
        clock.tick(60)
        a.req("POST", f"/sessions/{sid}/end", {})
        a.req("POST", f"/cards/{card}/advance", {"fromStationId": ids["Cut"], "toStationId": ids["Pack"],
                                                 "quantity": 10})
        sid = _start(a, ids, card, station="Pack", op="bob-jones")
        clock.tick(30)
        a.req("POST", f"/sessions/{sid}/end", {})
        a.req("POST", f"/cards/{card}/advance", {"fromStationId": ids["Pack"], "toStationId": ids["Store"],
                                                 "quantity": 10})
        clock.tick(1110)
    from app.routers import analysis as analysis_router
    monkeypatch.setattr(analysis_router, "utcnow", clock)
    api = R.Api(a.tc, a.company)
    data, questions = R.build_input(api, [{"name": "Pack twice as fast",
                                           "changes": [{"stationId": ids["Pack"], "speedFactor": 0.5}]}],
                                    now=clock.t)
    roles = {s["name"]: s["role"] for s in data["line"]["stations"]}
    assert roles == {"Cut": "work", "Pack": "work", "Store": "storage"}
    assert {s["station"] for s in data["statistics"]} >= {"Cut", "Pack"}
    ev = data["record_evidence"]["stations"]
    assert ev["Cut"]["workingShare"] > 0 and ev["Pack"]["workingShare"] > 0
    assert questions == ["Pack twice as fast"]
    assert data["what_if"]["fit_check"]["observedLeadTimeHours"] > 0
    text = R.prompt_text(data)
    for private in ("ann-smith", "bob-jones", ids["Cut"]):
        assert private not in text
    assert R.auto_scenarios(api.get("/sessions/stats"), {i: n for n, i in ids.items()}, {"Cut", "Pack"})[-1] \
        == {"name": "Demand +10%", "changes": [], "demandFactor": 1.1}


def test_the_export_names_each_sessions_lot(actor, clock):  # noqa: F811
    """The reviewer matches a session to its lot's moves by the lot id, which
    the export carries in the last column of Work sessions."""
    from openpyxl import load_workbook
    a = actor(); bootstrap_admin(a)
    ids, tmpl = _line(a)
    card = _lot(a, ids, tmpl)
    sid = _start(a, ids, card)
    clock.tick(60)
    a.req("POST", f"/sessions/{sid}/end", {})
    wb = load_workbook(io.BytesIO(a.req("GET", "/reports/export.xlsx").content), read_only=True)
    rows = list(wb["Work sessions"].iter_rows(values_only=True))
    assert rows[0][-1] == "Lot id" and rows[1][-1] == card


def test_nothing_is_sent_when_the_instructions_have_changed(monkeypatch):
    monkeypatch.setattr(R, "TASK", R.TASK + " Be generous.")
    monkeypatch.setattr("sys.argv", ["review_whatif.py", "--email", "x@example.com", "--auto"])
    called = []
    monkeypatch.setattr(R, "ask_claude", lambda *a, **k: called.append(1))
    with pytest.raises(SystemExit):
        R.main()
    assert not called
